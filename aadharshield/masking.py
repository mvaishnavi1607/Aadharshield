"""File-level masking: images (incl. multi-page TIFF), PDFs (scans + text layers) and Excel/CSV.

Every function takes bytes and returns (masked_bytes, report). Reports never contain full Aadhaar numbers
(only last-4 and positions)."""
import csv
import io
import re
from dataclasses import dataclass, field, asdict

import cv2
import numpy as np
from PIL import Image, ImageOps, ImageSequence

from . import detect, verhoeff
from .imgmask import mask_image, ImageResult

IMAGE_EXT = {".jpg", ".jpeg", ".png", ".tif", ".tiff", ".bmp", ".webp"}
PDF_EXT = {".pdf"}
SHEET_EXT = {".xlsx", ".xlsm", ".csv"}


@dataclass
class Report:
    kind: str
    numbers: int = 0
    qr: int = 0
    pages_hit: list = field(default_factory=list)
    needs_review: bool = False
    verified: bool = True
    notes: list = field(default_factory=list)
    last4: list = field(default_factory=list)

    def absorb(self, r: ImageResult, where=None):
        n = [f for f in r.findings if f.kind != "qr"]
        self.numbers += len(n)
        self.qr += sum(1 for f in r.findings if f.kind == "qr")
        self.last4 += [f.last4 for f in n]
        self.needs_review |= r.needs_review
        self.verified &= r.verified
        self.notes += [(f"{where}: " if where is not None else "") + x for x in r.notes]
        if r.findings and where is not None:
            self.pages_hit.append(where)

    def dict(self):
        return asdict(self)


def supported(ext: str) -> bool:
    return ext.lower() in IMAGE_EXT | PDF_EXT | SHEET_EXT


def mask_bytes(data: bytes, filename: str):
    ext = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext in IMAGE_EXT:
        return mask_image_file(data, ext)
    if ext in PDF_EXT:
        return mask_pdf(data)
    if ext == ".csv":
        return mask_csv(data)
    if ext in (".xlsx", ".xlsm"):
        return mask_xlsx(data, keep_vba=ext == ".xlsm")
    raise ValueError(f"unsupported file type {ext or '(none)'}")


# ---------------------------------------------------------------- images
def _to_bgr(pil: Image.Image):
    return cv2.cvtColor(np.array(pil.convert("RGB")), cv2.COLOR_RGB2BGR)


def mask_image_file(data: bytes, ext: str):
    im = Image.open(io.BytesIO(data))
    fmt = im.format or "PNG"
    rep = Report("image")
    frames = []
    for i, fr in enumerate(ImageSequence.Iterator(im)):
        fr = ImageOps.exif_transpose(fr.copy())
        res = mask_image(_to_bgr(fr))
        rep.absorb(res, i if getattr(im, "n_frames", 1) > 1 else None)
        frames.append(Image.fromarray(cv2.cvtColor(res.image, cv2.COLOR_BGR2RGB)))
    out = io.BytesIO()
    dpi = im.info.get("dpi")
    kw = {"dpi": dpi} if dpi else {}
    if fmt == "JPEG":
        frames[0].save(out, "JPEG", quality=92, **kw)          # re-encode: drops EXIF/GPS/thumbnail
    elif fmt == "TIFF" and len(frames) > 1:
        frames[0].save(out, "TIFF", save_all=True, append_images=frames[1:], compression="tiff_lzw", **kw)
    else:
        frames[0].save(out, fmt if fmt in ("PNG", "BMP", "WEBP", "TIFF") else "PNG", **kw)
    return out.getvalue(), rep


# ---------------------------------------------------------------- PDF
def mask_pdf(data: bytes):
    import pymupdf
    doc = pymupdf.open(stream=data, filetype="pdf")
    rep = Report("pdf")
    done = {}                      # xref -> ImageResult (images shared by several pages are masked once)
    for pno, page in enumerate(doc):
        hit = False
        # 1) text layer - work on glyph boxes so only the first 8 digits are redacted
        chars = [c for blk in page.get_text("rawdict")["blocks"] if blk["type"] == 0
                 for ln in blk["lines"] for sp in ln["spans"] for c in sp["chars"]]
        flat = "".join(c["c"] for c in chars)
        n_txt = 0
        for m in list(detect.AADHAAR_RE.finditer(flat)):
            d = m.group(1) + m.group(3) + m.group(5)
            if not verhoeff.is_aadhaar(d):
                continue
            idx = [i for i in range(m.start(), m.end()) if flat[i].isdigit()]
            c8, c9 = chars[idx[7]]["bbox"], chars[idx[8]]["bbox"]
            boxes = [chars[i]["bbox"] for i in range(idx[0], idx[7] + 1)]
            x1 = (c8[2] + c9[0]) / 2 if c9[0] > c8[2] else c8[2]
            r = pymupdf.Rect(min(b[0] for b in boxes), min(b[1] for b in boxes), x1, max(b[3] for b in boxes))
            page.add_redact_annot(r, fill=(0, 0, 0))
            n_txt += 1
            rep.last4.append(d[-4:])
        if n_txt:
            page.apply_redactions(images=pymupdf.PDF_REDACT_IMAGE_NONE)
            rep.numbers += n_txt
            hit = True
        # 2) embedded scans
        for info in page.get_images(full=True):
            xref = info[0]
            if xref not in done:
                done[xref] = None
                try:
                    raw = doc.extract_image(xref)["image"]
                    bgr = cv2.imdecode(np.frombuffer(raw, np.uint8), cv2.IMREAD_COLOR)
                except Exception:
                    bgr = None
                if bgr is None or min(bgr.shape[:2]) < 60:
                    continue
                res = mask_image(bgr)
                done[xref] = res
                if res.findings:
                    ok, buf = cv2.imencode(".jpg", res.image, [cv2.IMWRITE_JPEG_QUALITY, 92])
                    page.replace_image(xref, stream=buf.tobytes())
                rep.absorb(res, pno)
                hit |= bool(res.findings)
            elif done[xref] is not None and done[xref].findings:
                hit = True
                if pno not in rep.pages_hit:
                    rep.pages_hit.append(pno)
        if hit and pno not in rep.pages_hit:
            rep.pages_hit.append(pno)
    doc.set_metadata({})
    doc.del_xml_metadata()
    out = doc.tobytes(garbage=4, deflate=True, clean=True)
    return out, rep


# ---------------------------------------------------------------- sheets
def _cell_mask(v, force):
    """-> (new_value, count). Numeric 12-digit cells become masked text."""
    if isinstance(v, bool) or v is None:
        return v, 0
    if isinstance(v, (int, float)):
        if float(v).is_integer() and len(str(int(v))) in (12, 16):
            new, c = detect.mask_text(str(int(v)), force=force)
            return (new, c) if c else (v, 0)
        return v, 0
    if isinstance(v, str):
        new, c = detect.mask_text(v, force=force)
        return (new, c) if c else (v, 0)
    return v, 0


def mask_xlsx(data: bytes, keep_vba: bool = False):
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), keep_vba=keep_vba)
    rep = Report("xlsx")
    for ws in wb.worksheets:
        hot = set()
        for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 10)):   # header hint in first rows
            for c in row:
                if isinstance(c.value, str) and detect.HEADER_HINT.search(c.value):
                    hot.add(c.column)
        for row in ws.iter_rows():
            for c in row:
                if c.value is None:
                    continue
                new, n = _cell_mask(c.value, force=c.column in hot)
                if n:
                    if isinstance(c.value, (int, float)):
                        c.number_format = "@"
                    c.value = new
                    rep.numbers += n
                    rep.pages_hit.append(f"{ws.title}!{c.coordinate}")
                if c.comment and c.comment.text:
                    nt, n2 = detect.mask_text(c.comment.text, force=False)
                    if n2:
                        c.comment.text = nt
                        rep.numbers += n2
                if c.hyperlink and c.hyperlink.target:
                    nt, n3 = detect.mask_text(c.hyperlink.target)
                    if n3:
                        c.hyperlink.target = nt
                        rep.numbers += n3
        if hot:
            rep.notes.append(f"{ws.title}: Aadhaar-titled column(s) masked in force mode")
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue(), rep


def mask_csv(data: bytes):
    enc = "utf-8-sig"
    try:
        text = data.decode(enc)
    except UnicodeDecodeError:
        enc, text = "latin-1", data.decode("latin-1")
    rows = list(csv.reader(io.StringIO(text)))
    rep = Report("csv")
    hot = set()
    for r in rows[:10]:
        for j, v in enumerate(r):
            if detect.HEADER_HINT.search(v or ""):
                hot.add(j)
    for r in rows:
        for j, v in enumerate(r):
            new, n = detect.mask_text(v, force=j in hot)
            if n:
                r[j] = new
                rep.numbers += n
    out = io.StringIO()
    csv.writer(out, lineterminator="\r\n").writerows(rows)
    return out.getvalue().encode(enc), rep
